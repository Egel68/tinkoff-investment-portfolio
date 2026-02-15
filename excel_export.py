"""Красивый Excel-отчёт."""

from decimal import Decimal

from openpyxl import Workbook
from openpyxl.chart import PieChart as XlPieChart
from openpyxl.chart import Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.drawing.fill import ColorChoice, PatternFillProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from api_client import AccountInfo, AggregatedPosition, TypeAllocation
from config import COLORS


class ExcelExporter:
    def __init__(self):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)

        self.hdr_fill = PatternFill("solid", fgColor=COLORS["header_fill"])
        self.hdr_font = Font("Calibri", 11, bold=True, color=COLORS["header_font"])
        self.alt_fill = PatternFill("solid", fgColor=COLORS["alt_row"])
        self.thin_border = Border(
            left=Side("thin", color=COLORS["border"]),
            right=Side("thin", color=COLORS["border"]),
            top=Side("thin", color=COLORS["border"]),
            bottom=Side("thin", color=COLORS["border"]),
        )
        self.font_n = Font("Calibri", 10, color=COLORS["neutral"])
        self.font_b = Font("Calibri", 10, bold=True, color=COLORS["neutral"])
        self.font_title = Font("Calibri", 14, bold=True, color=COLORS["header_fill"])

    def _pnl_font(self, v: Decimal, bold=False) -> Font:
        c = (
            COLORS["positive"]
            if v > 0
            else COLORS["negative"]
            if v < 0
            else COLORS["neutral"]
        )
        return Font("Calibri", 10, bold=bold, color=c)

    def _style_header_row(self, ws, row, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row, c)
            cell.fill = self.hdr_fill
            cell.font = self.hdr_font
            cell.alignment = Alignment("center", "center", wrap_text=True)
            cell.border = self.thin_border

    def _auto_width(self, ws):
        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            max_length = 0
            for row_idx in range(1, ws.max_row + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                try:
                    if cell.value is not None:
                        max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[letter].width = min(max(max_length + 3, 10), 38)

    def export(
        self,
        accounts: list[AccountInfo],
        total_capital: Decimal,
        filepath: str,
        aggregated: list[AggregatedPosition] = None,
        allocations: list[TypeAllocation] = None,
    ) -> str:
        """Главный метод — экспорт в Excel."""
        # 1) Сводка
        ws0 = self.wb.create_sheet("Сводка")
        self._write_summary(ws0, accounts, total_capital)

        # 2) По типам
        if allocations:
            ws_alloc = self.wb.create_sheet("По типам")
            self._write_allocation(ws_alloc, allocations, total_capital)

        # 3) Все позиции
        if aggregated:
            ws_all = self.wb.create_sheet("Все позиции")
            self._write_aggregated(ws_all, aggregated, total_capital)

        # 4) По счетам
        for i, acc in enumerate(accounts):
            safe = (
                acc.name.replace("/", "-")
                .replace("\\", "-")
                .replace("*", "")
                .replace("?", "")
                .replace("[", "(")
                .replace("]", ")")
            )
            name = f"{i + 1}. {safe}"[:31]
            ws = self.wb.create_sheet(name)
            self._write_account(ws, acc, total_capital)

        self.wb.save(filepath)
        return filepath

    def _write_summary(self, ws, accounts, total_capital):
        ws.merge_cells("A1:G1")
        ws["A1"].value = "Сводка по всем счетам Т-Инвестиций"
        ws["A1"].font = self.font_title
        ws["A1"].alignment = Alignment("center")

        ws.merge_cells("A2:G2")
        ws["A2"].value = f"Общий капитал: {total_capital:,.2f} ₽"
        ws["A2"].font = Font("Calibri", 12, bold=True, color=COLORS["positive"])
        ws["A2"].alignment = Alignment("center")

        hdrs = [
            "Счёт",
            "Тип",
            "Статус",
            "Дата открытия",
            "Стоимость, ₽",
            "Доля от капитала, %",
            "Позиций",
        ]
        r = 4
        for c, h in enumerate(hdrs, 1):
            ws.cell(r, c, h)
        self._style_header_row(ws, r, len(hdrs))

        for acc in accounts:
            r += 1
            share = (
                (acc.total_value / total_capital * 100) if total_capital else Decimal(0)
            )
            vals = [
                acc.name,
                acc.acc_type,
                acc.status,
                acc.opened_date,
                float(acc.total_value),
                float(share),
                len(acc.positions),
            ]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(r, c, v)
                cell.font = self.font_n
                cell.border = self.thin_border
                cell.alignment = Alignment("center")
                if c == 5:
                    cell.number_format = "#,##0.00"
                if c == 6:
                    cell.number_format = "0.00"
            if (r - 4) % 2 == 0:
                for c in range(1, len(hdrs) + 1):
                    ws.cell(r, c).fill = self.alt_fill

        r += 2
        ws.cell(r, 1, "Денежные средства:").font = self.font_b
        for acc in accounts:
            if acc.total_currencies:
                r += 1
                ws.cell(r, 1, acc.name).font = self.font_b
                c = 2
                for cur, amt in acc.total_currencies.items():
                    ws.cell(r, c, f"{amt:,.2f} {cur}").font = self.font_n
                    c += 1

        self._auto_width(ws)

    def _write_allocation(
        self, ws, allocations: list[TypeAllocation], total_capital: Decimal
    ):
        """Лист с распределением по типам + круговая диаграмма."""
        ws.merge_cells("A1:H1")
        ws["A1"].value = "Распределение портфеля по типам инструментов"
        ws["A1"].font = self.font_title
        ws["A1"].alignment = Alignment("center")

        ws.merge_cells("A2:H2")
        ws["A2"].value = f"Общий капитал: {total_capital:,.2f} ₽"
        ws["A2"].font = Font("Calibri", 12, bold=True, color=COLORS["positive"])
        ws["A2"].alignment = Alignment("center")

        hdrs = [
            "Тип",
            "Кол-во позиций",
            "Ср. стоимость, ₽",
            "Рыночная стоимость, ₽",
            "P&L, ₽",
            "P&L, %",
            "Доля, %",
        ]
        hr = 4
        for c, h in enumerate(hdrs, 1):
            ws.cell(hr, c, h)
        self._style_header_row(ws, hr, len(hdrs))

        for idx, alloc in enumerate(allocations):
            r = hr + idx + 1
            vals = [
                alloc.type_name_ru,
                alloc.positions_count,
                float(alloc.total_avg_cost),
                float(alloc.total_market_cost),
                float(alloc.profit_loss),
                float(alloc.profit_loss_pct),
                float(alloc.share_of_total),
            ]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(r, c, v)
                cell.font = self.font_n
                cell.border = self.thin_border
                cell.alignment = Alignment("center")
                if c in (3, 4, 5):
                    cell.number_format = "#,##0.00"
                if c in (6, 7):
                    cell.number_format = "0.00"
                if c == 5:
                    cell.font = self._pnl_font(alloc.profit_loss)
                if c == 6:
                    cell.font = self._pnl_font(alloc.profit_loss_pct)

            if idx % 2 == 0:
                for c in range(1, len(hdrs) + 1):
                    ws.cell(r, c).fill = self.alt_fill

        # Итого
        tr = hr + len(allocations) + 1
        ws.cell(tr, 1, "ИТОГО").font = self.font_b
        ws.cell(tr, 1).border = self.thin_border

        t_cnt = sum(a.positions_count for a in allocations)
        t_avg = sum(a.total_avg_cost for a in allocations)
        t_mkt = sum(a.total_market_cost for a in allocations)
        t_pnl = t_mkt - t_avg
        t_pnl_pct = (t_pnl / t_avg * 100) if t_avg else Decimal(0)
        t_share = sum(a.share_of_total for a in allocations)

        for c, v in {
            2: t_cnt,
            3: float(t_avg),
            4: float(t_mkt),
            5: float(t_pnl),
            6: float(t_pnl_pct),
            7: float(t_share),
        }.items():
            cell = ws.cell(tr, c, v)
            cell.font = (
                self.font_b if c != 5 else self._pnl_font(Decimal(str(v)), bold=True)
            )
            cell.border = self.thin_border
            cell.alignment = Alignment("center")
            if c in (3, 4, 5):
                cell.number_format = "#,##0.00"
            if c in (6, 7):
                cell.number_format = "0.00"

        # === Круговая диаграмма ===
        chart = XlPieChart()
        chart.title = "Распределение по типам"
        chart.style = 10
        chart.width = 18
        chart.height = 14

        # Данные: типы (колонка 1) и доли (колонка 7)
        data_start = hr + 1
        data_end = hr + len(allocations)

        labels = Reference(ws, min_col=1, min_row=data_start, max_row=data_end)
        values = Reference(ws, min_col=7, min_row=data_start, max_row=data_end)

        chart.add_data(values, titles_from_data=False)
        chart.set_categories(labels)

        # Подписи
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        chart.dataLabels.showCatName = True
        chart.dataLabels.showVal = False

        # Раскраска сегментов
        if chart.series:
            series = chart.series[0]
            for i, alloc in enumerate(allocations):
                pt = DataPoint(idx=i)
                color_hex = alloc.color.lstrip("#")
                pt.graphicalProperties.solidFill = color_hex
                series.data_points.append(pt)

        # Размещаем диаграмму справа от таблицы
        ws.add_chart(chart, "I4")

        self._auto_width(ws)

    def _write_aggregated(
        self, ws, aggregated: list[AggregatedPosition], total_capital: Decimal
    ):
        ws.merge_cells("A1:N1")
        ws["A1"].value = "Все позиции — сводка по всем счетам"
        ws["A1"].font = self.font_title
        ws["A1"].alignment = Alignment("center")

        ws.merge_cells("A2:N2")
        ws["A2"].value = (
            f"Общий капитал: {total_capital:,.2f} ₽  |  "
            f"Уникальных бумаг: {len(aggregated)}"
        )
        ws["A2"].font = Font("Calibri", 12, bold=True, color=COLORS["positive"])
        ws["A2"].alignment = Alignment("center")

        hdrs = [
            "№",
            "Тикер",
            "Название",
            "Тип",
            "Сектор",
            "Валюта",
            "Кол-во\n(всего)",
            "Ср.взвеш.\nцена",
            "Текущая\nцена",
            "Ср. стоим.\n(всего)",
            "Рыночная\nстоим.",
            "P&L",
            "P&L, %",
            "Доля от\nкапитала, %",
            "Счета",
        ]
        hr = 4
        for c, h in enumerate(hdrs, 1):
            ws.cell(hr, c, h)
        self._style_header_row(ws, hr, len(hdrs))

        for idx, p in enumerate(aggregated, 1):
            r = hr + idx
            accounts_str = ", ".join(p.accounts)
            vals = [
                idx,
                p.ticker,
                p.name,
                self._type_ru(p.instrument_type),
                p.sector,
                p.currency.upper() if p.currency else "",
                float(p.total_quantity),
                float(p.weighted_avg_price),
                float(p.current_price),
                float(p.total_avg_cost),
                float(p.total_market_cost),
                float(p.profit_loss),
                float(p.profit_loss_pct),
                float(p.share_of_total),
                accounts_str,
            ]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(r, c, v)
                cell.border = self.thin_border
                cell.alignment = Alignment("center", "center")
                cell.font = self.font_n
                if c in (8, 9, 10, 11, 12):
                    cell.number_format = "#,##0.00"
                if c in (13, 14):
                    cell.number_format = "0.00"
                if c == 12:
                    cell.font = self._pnl_font(p.profit_loss)
                if c == 13:
                    cell.font = self._pnl_font(p.profit_loss_pct)
                if c == 15:
                    cell.alignment = Alignment("left", "center", wrap_text=True)

            if idx % 2 == 0:
                for c in range(1, len(hdrs) + 1):
                    ws.cell(r, c).fill = self.alt_fill

        tr = hr + len(aggregated) + 1
        ws.cell(tr, 1, "ИТОГО").font = self.font_b
        ws.cell(tr, 1).border = self.thin_border

        t_avg = sum(p.total_avg_cost for p in aggregated)
        t_mkt = sum(p.total_market_cost for p in aggregated)
        t_pnl = t_mkt - t_avg
        t_pnl_pct = (t_pnl / t_avg * 100) if t_avg else Decimal(0)
        t_share = sum(p.share_of_total for p in aggregated)

        for c, v in {
            10: t_avg,
            11: t_mkt,
            12: t_pnl,
            13: t_pnl_pct,
            14: t_share,
        }.items():
            cell = ws.cell(tr, c, float(v))
            cell.font = self.font_b if c != 12 else self._pnl_font(v, bold=True)
            cell.border = self.thin_border
            cell.alignment = Alignment("center")
            cell.number_format = "#,##0.00" if c in (10, 11, 12) else "0.00"

        ws.freeze_panes = f"A{hr + 1}"
        self._auto_width(ws)

    def _write_account(self, ws, acc: AccountInfo, total_capital: Decimal):
        ws.merge_cells("A1:P1")
        ws["A1"].value = f"{acc.name} ({acc.acc_type})"
        ws["A1"].font = self.font_title
        ws["A1"].alignment = Alignment("center")

        ws.merge_cells("A2:H2")
        ws["A2"].value = (
            f"Статус: {acc.status} | Открыт: {acc.opened_date} | "
            f"Стоимость: {acc.total_value:,.2f} ₽"
        )
        ws["A2"].font = Font("Calibri", 11, color=COLORS["neutral"])

        if acc.total_currencies:
            ws.merge_cells("A3:H3")
            ws["A3"].value = "Деньги: " + " | ".join(
                f"{cur}: {amt:,.2f}" for cur, amt in acc.total_currencies.items()
            )
            ws["A3"].font = Font("Calibri", 10, italic=True)

        hdrs = [
            "№",
            "Тикер",
            "Название",
            "Тип",
            "Сектор",
            "Страна",
            "Валюта",
            "Кол-во",
            "Ср. цена\nпокупки",
            "Текущая\nцена",
            "Ср. стоимость\nпозиции",
            "Рыночная\nстоимость",
            "P&L",
            "P&L, %",
            "Доля от\nсчёта, %",
            "Доля от\nкапитала, %",
        ]
        hr = 5
        for c, h in enumerate(hdrs, 1):
            ws.cell(hr, c, h)
        self._style_header_row(ws, hr, len(hdrs))

        for idx, p in enumerate(acc.positions, 1):
            r = hr + idx
            vals = [
                idx,
                p.ticker,
                p.name,
                self._type_ru(p.instrument_type),
                p.sector,
                p.country,
                p.currency.upper() if p.currency else "",
                float(p.quantity),
                float(p.average_price),
                float(p.current_price),
                float(p.average_cost),
                float(p.market_cost),
                float(p.profit_loss),
                float(p.profit_loss_pct),
                float(p.share_of_account),
                float(p.share_of_total),
            ]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(r, c, v)
                cell.border = self.thin_border
                cell.alignment = Alignment("center", "center")
                cell.font = self.font_n
                if c in (9, 10, 11, 12, 13):
                    cell.number_format = "#,##0.00"
                if c in (14, 15, 16):
                    cell.number_format = "0.00"
                if c == 13:
                    cell.font = self._pnl_font(p.profit_loss)
                if c == 14:
                    cell.font = self._pnl_font(p.profit_loss_pct)

            if idx % 2 == 0:
                for c in range(1, len(hdrs) + 1):
                    ws.cell(r, c).fill = self.alt_fill

        tr = hr + len(acc.positions) + 1
        ws.cell(tr, 1, "ИТОГО").font = self.font_b
        ws.cell(tr, 1).border = self.thin_border

        t_avg = sum(p.average_cost for p in acc.positions)
        t_mkt = sum(p.market_cost for p in acc.positions)
        t_pnl = t_mkt - t_avg
        t_pnl_pct = (t_pnl / t_avg * 100) if t_avg else Decimal(0)
        t_sh_a = sum(p.share_of_account for p in acc.positions)
        t_sh_c = sum(p.share_of_total for p in acc.positions)

        for c, v in {
            11: t_avg,
            12: t_mkt,
            13: t_pnl,
            14: t_pnl_pct,
            15: t_sh_a,
            16: t_sh_c,
        }.items():
            cell = ws.cell(tr, c, float(v))
            cell.font = self.font_b if c != 13 else self._pnl_font(v, bold=True)
            cell.border = self.thin_border
            cell.alignment = Alignment("center")
            cell.number_format = "#,##0.00" if c in (11, 12, 13) else "0.00"

        ws.freeze_panes = f"A{hr + 1}"
        self._auto_width(ws)

    @staticmethod
    def _type_ru(t: str) -> str:
        return {
            "share": "Акция",
            "bond": "Облигация",
            "etf": "Фонд (ETF)",
            "currency": "Валюта",
            "futures": "Фьючерс",
            "option": "Опцион",
            "sp": "Структ. продукт",
        }.get(t, t)
